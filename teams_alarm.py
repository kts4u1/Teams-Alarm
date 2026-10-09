"""
Teams 'MR Schedule' 채팅에서 내 이름 줄의 'PM' 병원을 찾아, 그 병원의 GA/GC 압력을
Teams 나에게 메시지로 보내는 프로그램.

흐름
----
1. Edge로 Teams 웹을 열어 MR Schedule 채팅 내용을 읽는다 (처음 한 번만 로그인).
2. 오늘 날짜 아래 '김태수 : 병원 PM' 줄에서 PM 병원 이름을 뽑는다.
3. 헬륨 엑셀의 SRN/Site 표로 병원 이름 -> 시스템 번호(SRN)를 찾는다.
4. RADAR에서 시스템별 GA/GC 최신 값을 가져온다 (처음 한 번만 로그인).
5. 결과를 Windows 알림창으로 띄운다.
   (TEAMS_WEBHOOK_URL 환경변수가 있으면 Teams 웹훅으로 나에게 메시지를 보낸다.)

사전 준비
---------
pip install selenium webdriver-manager openpyxl
set TEAMS_WEBHOOK_URL=<Workflows 웹훅 주소>         (선택. 없으면 Windows 알림창)
set HELIUM_XLSX=C:\\경로\\01. 2026 Q4_헬륨 파일_여기입력.xlsx   (선택)

실행
----
py teams_alarm.py            # 실제 실행
py teams_alarm.py --dry-run  # 알림을 보내지 않고 콘솔에만 출력
py teams_alarm.py --hospital 송도지안   # Teams를 읽지 않고 지정한 병원으로 테스트
"""

import glob
import json
import os
import re
import sys
import ctypes
import time
import urllib.request
from datetime import datetime

import openpyxl

import radar_gradient_pressure as radar

MY_NAME = os.environ.get("ALARM_NAME", "김태수")
CHAT_NAME = "MR Schedule"
TEAMS_URL = "https://teams.microsoft.com/v2/"
LOGIN_WAIT_TIMEOUT = 180
PROFILE_DIR = os.path.join(
    os.environ.get("LOCALAPPDATA", os.path.expanduser("~")), "RadarExport", "edge-profile"
)
ONEDRIVE_DIRS = [d for d in {os.environ.get("OneDriveCommercial"), os.environ.get("OneDrive"),
                                   *glob.glob(os.path.join(os.path.expanduser("~"), "OneDrive*"))} if d]

OPEN_CHAT_JS = """
function findByText(root, text) {
    for (const el of root.querySelectorAll('*')) {
        if (el.shadowRoot) {
            const f = findByText(el.shadowRoot, text);
            if (f) return f;
        }
        if (el.children.length === 0 && el.textContent && el.textContent.trim() === text) return el;
    }
    return null;
}
const t = findByText(document, arguments[0]);
if (t) { t.click(); return true; }
return false;
"""


# ---------- Teams 메시지 해석 ----------

def today_header(now=None):
    now = now or datetime.now()
    return f"{now.year}년 {now.month}월 {now.day}일"


def pm_hospitals(text, name=MY_NAME, now=None):
    """채팅 전체 텍스트에서 오늘 날짜 아래 `name` 줄의 PM 병원 이름 목록을 돌려줍니다."""
    header = today_header(now)
    start = text.rfind(header)
    if start < 0:
        return None  # 오늘 일정이 아직 올라오지 않음
    section = text[start + len(header):]
    next_date = re.search(r"\d{4}년 \d{1,2}월 \d{1,2}일", section)
    if next_date:
        section = section[:next_date.start()]

    hospitals = []
    for m in re.finditer(rf"^\s*\d+\s*\.\s*{re.escape(name)}\s*[:：]?[ \t]*(.*)$", section, re.M):
        for part in re.split(r"[,/、&+]|및", m.group(1)):
            if not re.search(r"\bPM\b", part, re.I):
                continue
            hospital = re.sub(r"\bPM\b", "", part, flags=re.I).strip(" -·()")
            if hospital:
                hospitals.append(hospital)
    return hospitals


# ---------- 병원 이름 -> SRN ----------

def normalize(name):
    name = re.sub(r"\([^)]*\)", "", str(name))
    return re.sub(r"\s+", "", name).lower()


def find_helium_xlsx():
    path = os.environ.get("HELIUM_XLSX")
    if path:
        return path.strip('"')
    candidates = []
    for root in ONEDRIVE_DIRS:
        candidates += glob.glob(os.path.join(root, "**", "01.*헬륨*.xlsx"), recursive=True)
    candidates = [c for c in candidates if not os.path.basename(c).startswith("~$")]
    if not candidates:
        raise FileNotFoundError(
            "헬륨 엑셀을 찾지 못했습니다. PowerShell에서 아래처럼 경로를 지정하세요.\n"
            '$env:HELIUM_XLSX = "C:\\전체\\경로\\파일.xlsx"')
    return max(candidates, key=os.path.getmtime)


def load_sites(path):
    """첫 번째(최신 분기) 시트의 (SRN, Site) 목록을 읽습니다."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c else "" for c in next(rows)]
    srn_col, site_col = header.index("SRN"), header.index("Site")
    return [(str(r[srn_col]).strip(), str(r[site_col]).strip())
            for r in rows if r[srn_col] and r[site_col]]


def find_srns(hospital, sites):
    """병원 이름과 맞는 (SRN, Site) 목록. 정확히 같은 이름이 있으면 그것만, 없으면 포함 관계로 찾습니다."""
    key = normalize(hospital)
    exact = [(s, n) for s, n in sites if normalize(n) == key]
    if exact:
        return exact
    if len(key) < 2:
        return []
    return [(s, n) for s, n in sites if key in normalize(n) or normalize(n) in key]


# ---------- 알림 메시지 ----------

def format_value(item):
    if item is None:
        return "데이터 없음"
    mark = {"OK": "✅", "LOW": "⚠️ 하한 미만", "HIGH": "⚠️ 상한 초과"}[item["status"]]
    return f"{item['value']:.2f} bar {mark}"


def build_message(results, missing, now=None):
    now = now or datetime.now()
    lines = [f"[GA/GC 압력] {now:%Y-%m-%d} PM 병원"]
    for hospital, systems in results.items():
        for srn, site, values in systems:
            if values.get("skip"):  # RADAR 목록에 없는 시스템은 알리지 않음
                continue
            if "error" in values:
                lines.append(f"{site} (SRN {srn}) - 조회 실패: {values['error']}")
                continue
            date = (values["GA"] or values["GC"] or {}).get("date", "")
            lines.append(f"{site} (SRN {srn}) - GA {format_value(values['GA'])} / "
                         f"GC {format_value(values['GC'])}" + (f" [{date} 기준]" if date else ""))
    for hospital in missing:
        lines.append(f"{hospital}: 헬륨 엑셀에서 병원을 찾지 못했습니다")
    return "\n".join(lines)


def send_to_teams(webhook_url, text):
    card = {
        "type": "message",
        "attachments": [{
            "contentType": "application/vnd.microsoft.card.adaptive",
            "content": {
                "type": "AdaptiveCard", "version": "1.4",
                "$schema": "http://adaptivecards.io/schemas/adaptive-card.json",
                "body": [{"type": "TextBlock", "text": text, "wrap": True}],
            },
        }],
    }
    req = urllib.request.Request(webhook_url, data=json.dumps(card).encode("utf-8"),
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        return resp.status


def show_popup(text, title="GA/GC 압력 알림"):
    # MB_ICONINFORMATION | MB_SETFOREGROUND | MB_TOPMOST
    ctypes.windll.user32.MessageBoxW(0, text, title, 0x40 | 0x10000 | 0x40000)


# ---------- 브라우저 ----------

def read_chat_text(driver):
    driver.get(TEAMS_URL)
    end_time = time.time() + LOGIN_WAIT_TIMEOUT
    header = today_header()
    opened = False
    while time.time() < end_time:
        try:
            if not opened:
                opened = driver.execute_script(OPEN_CHAT_JS, CHAT_NAME)
            else:
                text = driver.execute_script("return document.body.innerText")
                if header in text:
                    return text
        except Exception:
            pass
        time.sleep(2)
    return driver.execute_script("return document.body.innerText") if opened else None


def main():
    dry_run = "--dry-run" in sys.argv
    webhook = os.environ.get("TEAMS_WEBHOOK_URL")

    os.makedirs(PROFILE_DIR, exist_ok=True)
    driver = radar.setup_driver(PROFILE_DIR)
    try:
        if "--hospital" in sys.argv:
            hospitals = [sys.argv[sys.argv.index("--hospital") + 1]]
        else:
            text = read_chat_text(driver)
            if text is None:
                sys.exit("Teams에서 MR Schedule 채팅을 찾지 못했습니다.")
            hospitals = pm_hospitals(text)
            if hospitals is None:
                sys.exit("오늘 날짜의 MR Schedule 메시지가 아직 없습니다.")
        if not hospitals:
            print("오늘은 PM 병원이 없습니다.")
            return

        sites = load_sites(find_helium_xlsx())
        matched = {h: find_srns(h, sites) for h in hospitals}
        missing = [h for h, found in matched.items() if not found]
        srns = sorted({s for found in matched.values() for s, _ in found})

        pressures = radar.get_gradient_pressure(srns, driver) if srns else {}
        results = {h: [(s, n, pressures[s]) for s, n in found] for h, found in matched.items() if found}
    finally:
        driver.quit()

    message = build_message(results, missing)
    print(message)
    if dry_run:
        return
    if webhook:
        print("Teams 전송 결과:", send_to_teams(webhook, message))
    else:
        show_popup(message)


if __name__ == "__main__":
    main()
